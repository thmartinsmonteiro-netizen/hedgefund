"""Rebuild reports/tracker.xlsx from the CSVs and the NAV engine output. Values only:
the arithmetic lives in nav.py, the workbook is a view for reading on a phone."""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter as L

ROOT = Path(__file__).resolve().parents[1]
DATA, REPORTS = ROOT / "data", ROOT / "reports"
BOLD = Font(name="Arial", bold=True); NORM = Font(name="Arial"); GREY = PatternFill("solid", fgColor="DDDDDD")
FMT = {"usd": '$#,##0;($#,##0);-', "pct": '0.00%;(0.00%);-', "px": '#,##0.00'}

def sheet(wb, name, df, fmts=None, widths=None):
    ws = wb.create_sheet(name)
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=j, value=c); cell.font = BOLD; cell.fill = GREY
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, row in enumerate(df.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if pd.isna(v) if not isinstance(v, str) else False:
                v = None
            if hasattr(v, "isoformat") and not isinstance(v, str):
                v = v.isoformat()[:10]
            cell = ws.cell(row=i, column=j, value=v); cell.font = NORM
            f = (fmts or {}).get(df.columns[j-1])
            if f: cell.number_format = FMT[f]
    for j, c in enumerate(df.columns, 1):
        ws.column_dimensions[L(j)].width = (widths or {}).get(c, max(10, min(40, len(str(c)) + 4)))
    ws.freeze_panes = "B2"
    return ws

def main():
    wb = Workbook(); wb.remove(wb.active)
    nav = pd.read_csv(REPORTS / "nav.csv")
    sheet(wb, "NAV", nav, {"cash_usd": "usd", "positions_mv_usd": "usd", "interest_usd": "usd", "net_flows_usd": "usd",
                           "nav_usd": "usd", "core_bench_cum": "pct", "cash_bench_cum": "pct", "daily_return": "pct",
                           "cum_return": "pct", "drawdown": "pct", "alpha_vs_core": "pct", "alpha_vs_cash": "pct"}, {"flags": 40})
    sheet(wb, "Positions", pd.read_csv(REPORTS / "positions_latest.csv"),
          {"cost_basis_usd": "usd", "mv_usd": "usd", "pnl_usd": "usd", "pnl_pct": "pct", "pct_nav": "pct",
           "upside_to_target": "pct", "last_close": "px", "target_price": "px"}, {"name": 28, "kill_criteria": 50})
    sheet(wb, "Trades", pd.read_csv(DATA / "trades.csv"), {"price_native": "px"}, {"thesis": 60, "notes": 50})
    sheet(wb, "CashFlows", pd.read_csv(DATA / "cashflows.csv"), {"amount_usd": "usd"}, {"note": 50})
    sheet(wb, "Pipeline", pd.read_csv(DATA / "pipeline.csv"), {"price_at_log": "px", "target": "px"}, {"catalyst": 40, "kill_criteria": 50})
    sheet(wb, "Rejected", pd.read_csv(DATA / "rejected.csv"), None, {"why_passed": 70, "reconsider_if": 50})
    sheet(wb, "Closed", pd.read_csv(DATA / "closed.csv"), {"cost_basis_usd": "usd", "proceeds_usd": "usd"}, {"why_exited": 40, "what_it_taught": 50})
    sheet(wb, "Prices", pd.read_csv(DATA / "prices.csv"), {"close": "px"})
    wb.save(REPORTS / "tracker.xlsx"); print("wrote", REPORTS / "tracker.xlsx")

if __name__ == "__main__":
    main()
